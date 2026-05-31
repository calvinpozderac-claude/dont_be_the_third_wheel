#!/usr/bin/env python3
"""Generate Three's a Crowd print-and-play card sheets (4 pages, duplex-ready)."""

from docx import Document
from docx.shared import Inches, Pt, RGBColor, Twips
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

# ── Card data ──────────────────────────────────────────────────────────────────

CARD_ACTIONS = {
    1: "Switch any two date piles",
    2: "Swap arrival order of any two cards",
    3: "Deactivate a bonus or activate a debuff",
    4: "Switch any two date piles",
    5: "Swap arrival order of any two cards",
    6: "Deactivate a bonus or activate a debuff",
}
CARD_MODIFIERS = {
    1: "+1 point at end of game",
    2: "Your score ×2 this date",
    3: "Peek at one player's move",
    4: "−1 point at end of game",
    5: "Your score ÷2 this date",
    6: "Show your move to one player",
}

CARDS = [
    # ── Red player ─────────────────────────────────────────────────────────────
    {"name": "Funny Jokes",          "type": "bonus",  "num": 1, "color": "red"},
    {"name": "Good Listener",        "type": "bonus",  "num": 2, "color": "red"},
    {"name": "Paid for the Date",    "type": "bonus",  "num": 3, "color": "red"},
    {"name": "On Their Phone",       "type": "debuff", "num": 4, "color": "red"},
    {"name": "Unkempt",              "type": "debuff", "num": 5, "color": "red"},
    {"name": "Bad Breath",           "type": "debuff", "num": 6, "color": "red"},
    # ── Blue player ────────────────────────────────────────────────────────────
    {"name": "Well-Groomed",         "type": "bonus",  "num": 1, "color": "blue"},
    {"name": "Kind to Staff",        "type": "bonus",  "num": 2, "color": "blue"},
    {"name": "Thoughtful Questions", "type": "bonus",  "num": 3, "color": "blue"},
    {"name": "Constantly Complaining","type": "debuff","num": 4, "color": "blue"},
    {"name": "Self-Absorbed",        "type": "debuff", "num": 5, "color": "blue"},
    {"name": "Made Them Pay",        "type": "debuff", "num": 6, "color": "blue"},
    # ── Green player ───────────────────────────────────────────────────────────
    {"name": "Confident",            "type": "bonus",  "num": 1, "color": "green"},
    {"name": "Positive Attitude",    "type": "bonus",  "num": 2, "color": "green"},
    {"name": "Put Their Phone Away", "type": "bonus",  "num": 3, "color": "green"},
    {"name": "Talks About Exes",     "type": "debuff", "num": 4, "color": "green"},
    {"name": "Condescending",        "type": "debuff", "num": 5, "color": "green"},
    {"name": "Passing Gas",          "type": "debuff", "num": 6, "color": "green"},
]

# ── Palette ────────────────────────────────────────────────────────────────────

PALETTE = {
    "red":   {"bg": "FFCCCC", "border": "C03030", "text": "7A1010"},
    "blue":  {"bg": "CCE0FF", "border": "3050C0", "text": "1030A0"},
    "green": {"bg": "C8FFDC", "border": "20A040", "text": "0A6020"},
}

ACTION_ON_BG    = "FFF0A0"   # warm yellow  – action is active
ACTION_OFF_BG   = "EEEEEE"   # light grey   – action skipped
ACTION_ON_TEXT  = "5A4000"
ACTION_OFF_TEXT = "999999"

BONUS_ON_BG     = "B8FFD0"   # green        – bonus active
BONUS_ON_TEXT   = "0A5020"
BONUS_OFF_BG    = "EEEEEE"   # grey         – bonus forfeited
BONUS_OFF_TEXT  = "999999"

DEBUFF_ON_BG    = "FFD0CC"   # red          – debuff active
DEBUFF_ON_TEXT  = "7A1010"
DEBUFF_OFF_BG   = "EEEEEE"   # grey         – debuff avoided
DEBUFF_OFF_TEXT = "999999"

MM_BG   = "D0E4FF"
MM_TEXT = "153080"
PS_BG   = "D0F0D8"
PS_TEXT = "0A5020"

# ── XML helpers ────────────────────────────────────────────────────────────────

def _rgb(hex6):
    return RGBColor(int(hex6[0:2], 16), int(hex6[2:4], 16), int(hex6[4:6], 16))

def _shd(el, fill):
    s = OxmlElement("w:shd")
    s.set(qn("w:val"),   "clear")
    s.set(qn("w:color"), "auto")
    s.set(qn("w:fill"),  fill)
    el.append(s)

def set_cell_bg(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    _shd(tcPr, fill)

def set_cell_borders(cell, color, thick=24):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"),   "single")
        b.set(qn("w:sz"),    str(thick))
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), color)
        borders.append(b)
    tcPr.append(borders)

def set_cell_margins(cell, top=36, left=72, bottom=36, right=72):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:w"),    str(val))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tcPr.append(mar)

def set_row_height(row, inches):
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:trHeight")
    h.set(qn("w:val"),   str(int(Inches(inches) / 914.4 * 1440)))
    h.set(qn("w:hRule"), "atLeast")
    trPr.append(h)

def set_table_no_borders(table):
    tblPr = table._tbl.tblPr
    tb = OxmlElement("w:tblBorders")
    for edge in ("top","left","bottom","right","insideH","insideV"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "none")
        tb.append(b)
    tblPr.append(tb)

def zero_para_spacing(para):
    pf = para.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after  = Pt(0)
    pf.line_spacing = Pt(11)

# ── Card paragraph builders ────────────────────────────────────────────────────

def _para(cell, text, *, bold=False, italic=False, size=9, txt_color=None,
          bg=None, align=WD_ALIGN_PARAGRAPH.LEFT,
          sb=0, sa=1, line=None):
    p = cell.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(sb)
    pf.space_after  = Pt(sa)
    pf.line_spacing = Pt(line or size + 1.5)
    r = p.add_run(text)
    r.bold   = bold
    r.italic = italic
    r.font.size = Pt(size)
    if txt_color:
        r.font.color.rgb = _rgb(txt_color)
    if bg:
        rPr = r._r.get_or_add_rPr()
        _shd(rPr, bg)
    return p

def _divider(cell):
    _para(cell, "", size=3, sa=0, sb=0)

# ── Card face renderers ────────────────────────────────────────────────────────

def fill_front(cell, card):
    """Action side — Make a Move."""
    pal = PALETTE[card["color"]]
    is_bonus = card["type"] == "bonus"
    action   = CARD_ACTIONS[card["num"]]
    modifier = CARD_MODIFIERS[card["num"]]

    # clear default empty paragraph
    for p in list(cell.paragraphs):
        p._element.getparent().remove(p._element)

    set_cell_bg(cell, pal["bg"])
    set_cell_borders(cell, pal["border"], thick=24)
    set_cell_margins(cell, top=60, left=80, bottom=36, right=80)

    # ── Title ──
    _para(cell, card["name"], bold=True, size=11, txt_color=pal["text"],
          align=WD_ALIGN_PARAGRAPH.CENTER, sb=0, sa=2, line=13)

    # ── Type badge ──
    if is_bonus:
        _para(cell, "  BONUS  ", bold=True, size=7, txt_color=BONUS_ON_TEXT,
              bg=BONUS_ON_BG, align=WD_ALIGN_PARAGRAPH.CENTER, sa=4)
    else:
        _para(cell, "  DEBUFF  ", bold=True, size=7, txt_color=DEBUFF_ON_TEXT,
              bg=DEBUFF_ON_BG, align=WD_ALIGN_PARAGRAPH.CENTER, sa=4)

    # ── ACTION block (highlighted — being used) ──
    _para(cell, "⚡  ACTION", bold=True, size=7, txt_color=ACTION_ON_TEXT, sa=1)
    _para(cell, action, size=8, txt_color=ACTION_ON_TEXT,
          bg=ACTION_ON_BG, sa=5)

    # ── MODIFIER block ──
    if is_bonus:
        # Bonus forfeited
        _para(cell, "○  BONUS", bold=True, size=7, txt_color=BONUS_OFF_TEXT, sa=1)
        _para(cell, modifier, size=8, txt_color=BONUS_OFF_TEXT,
              bg=BONUS_OFF_BG, sa=0)
    else:
        # Debuff activated
        _para(cell, "✗  DEBUFF", bold=True, size=7, txt_color=DEBUFF_ON_TEXT, sa=1)
        _para(cell, modifier, size=8, txt_color=DEBUFF_ON_TEXT,
              bg=DEBUFF_ON_BG, sa=0)

    # ── spacer + MM bar ──
    _para(cell, "", size=4, sa=0)
    _para(cell, "  💘  Make a Move!  ", bold=True, size=9,
          txt_color=MM_TEXT, bg=MM_BG,
          align=WD_ALIGN_PARAGRAPH.CENTER, sb=4, sa=0)


def fill_back(cell, card):
    """Modifier side — Play it Safe."""
    pal = PALETTE[card["color"]]
    is_bonus = card["type"] == "bonus"
    action   = CARD_ACTIONS[card["num"]]
    modifier = CARD_MODIFIERS[card["num"]]

    for p in list(cell.paragraphs):
        p._element.getparent().remove(p._element)

    set_cell_bg(cell, pal["bg"])
    set_cell_borders(cell, pal["border"], thick=24)
    set_cell_margins(cell, top=60, left=80, bottom=36, right=80)

    # ── Title ──
    _para(cell, card["name"], bold=True, size=11, txt_color=pal["text"],
          align=WD_ALIGN_PARAGRAPH.CENTER, sb=0, sa=2, line=13)

    # ── Type badge ──
    if is_bonus:
        _para(cell, "  BONUS  ", bold=True, size=7, txt_color=BONUS_ON_TEXT,
              bg=BONUS_ON_BG, align=WD_ALIGN_PARAGRAPH.CENTER, sa=4)
    else:
        _para(cell, "  DEBUFF  ", bold=True, size=7, txt_color=DEBUFF_ON_TEXT,
              bg=DEBUFF_ON_BG, align=WD_ALIGN_PARAGRAPH.CENTER, sa=4)

    # ── ACTION block (greyed — skipped) ──
    _para(cell, "○  ACTION", bold=True, size=7, txt_color=ACTION_OFF_TEXT, sa=1)
    _para(cell, action, size=8, txt_color=ACTION_OFF_TEXT,
          bg=ACTION_OFF_BG, sa=5)

    # ── MODIFIER block ──
    if is_bonus:
        # Bonus active!
        _para(cell, "✓  BONUS", bold=True, size=7, txt_color=BONUS_ON_TEXT, sa=1)
        _para(cell, modifier, size=8, txt_color=BONUS_ON_TEXT,
              bg=BONUS_ON_BG, sa=0)
    else:
        # Debuff avoided
        _para(cell, "○  DEBUFF", bold=True, size=7, txt_color=DEBUFF_OFF_TEXT, sa=1)
        _para(cell, modifier, size=8, txt_color=DEBUFF_OFF_TEXT,
              bg=DEBUFF_OFF_BG, sa=0)

    # ── spacer + PS bar ──
    _para(cell, "", size=4, sa=0)
    _para(cell, "  🛡  Play it Safe  ", bold=True, size=9,
          txt_color=PS_TEXT, bg=PS_BG,
          align=WD_ALIGN_PARAGRAPH.CENTER, sb=4, sa=0)


# ── Page builder ───────────────────────────────────────────────────────────────

def add_card_page(doc, nine_cards, fill_fn):
    tbl = doc.add_table(rows=3, cols=3)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_no_borders(tbl)

    # Fixed column widths
    for row in tbl.rows:
        set_row_height(row, 3.5)
        for ci, cell in enumerate(row.cells):
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            tcW = OxmlElement("w:tcW")
            tcW.set(qn("w:w"),    str(int(Inches(2.5) / 914.4 * 1440)))
            tcW.set(qn("w:type"), "dxa")
            tcPr.append(tcW)

    for ri, row in enumerate(tbl.rows):
        for ci, cell in enumerate(row.cells):
            idx = ri * 3 + ci
            if idx < len(nine_cards):
                fill_fn(cell, nine_cards[idx])
            else:
                set_cell_borders(cell, "CCCCCC", thick=4)


def mirror_rows(cards_9):
    """Reverse each row of 3 for duplex alignment (flip on long edge)."""
    out = []
    for i in range(0, 9, 3):
        out.extend(reversed(cards_9[i:i+3]))
    return out


# ── Main ───────────────────────────────────────────────────────────────────────

def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width      = Inches(8.5)
    sec.page_height     = Inches(11)
    sec.left_margin     = Inches(0.5)
    sec.right_margin    = Inches(0.5)
    sec.top_margin      = Inches(0.25)
    sec.bottom_margin   = Inches(0.25)

    # Remove default Normal style spacing
    normal = doc.styles["Normal"]
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after  = Pt(0)

    sheet1 = CARDS[:9]
    sheet2 = CARDS[9:]

    # Sheet 1 — fronts
    add_card_page(doc, sheet1, fill_front)
    doc.add_page_break()

    # Sheet 1 — backs (mirror columns for duplex)
    add_card_page(doc, mirror_rows(sheet1), fill_back)
    doc.add_page_break()

    # Sheet 2 — fronts
    add_card_page(doc, sheet2, fill_front)
    doc.add_page_break()

    # Sheet 2 — backs (mirror)
    add_card_page(doc, mirror_rows(sheet2), fill_back)

    out = "/home/user/dont_be_the_third_wheel/threes_a_crowd_cards.docx"
    doc.save(out)
    print(f"Saved → {out}")


if __name__ == "__main__":
    build()
